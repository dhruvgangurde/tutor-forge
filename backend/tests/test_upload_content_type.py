"""
tests/test_upload_content_type.py
---------------------------------
Upload type checks look at the bytes, not just the name (audit 2026-10-06 #7a).

A file is accepted only when its extension, its declared MIME type and its
leading bytes all agree. Mismatches are a 400 with a message naming the file,
raised before any course or ingestion job is created.
"""

import io
import zipfile

import pytest
from httpx import ASGITransport, AsyncClient
from pptx import Presentation
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from core.database import Base
from core.dependencies import get_db_session
from core.security import hash_password
from courses.file_types import content_mismatch
from db.models import Course, User
from main import app

PDF = "application/pdf"
PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
PPT = "application/vnd.ms-powerpoint"  # legacy type, no longer accepted
TXT = "text/plain"

PDF_BYTES = b"%PDF-1.7\n1 0 obj\n<< /Type /Catalog >>\nendobj\n%%EOF\n"
PPT_BYTES = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 504
EXE_BYTES = b"MZ\x90\x00\x03\x00\x00\x00\x04\x00\x00\x00\xff\xff" + b"\x00" * 64


def _pptx_bytes() -> bytes:
    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[0])
    buf = io.BytesIO()
    deck.save(buf)
    return buf.getvalue()


def _plain_zip() -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("readme.md", "not a presentation")
    return buf.getvalue()


# ── The check itself ──────────────────────────────────────────────────────────

@pytest.mark.parametrize(
    "name, ext, mime, content",
    [
        ("notes.pdf", ".pdf", PDF, PDF_BYTES),
        ("slides.pptx", ".pptx", PPTX, _pptx_bytes()),
        ("notes.txt", ".txt", TXT, "Plain notes, with UTF-8: café.\n".encode()),
        ("notes.txt", ".txt", "text/plain; charset=utf-8", b"notes"),
        ("empty.txt", ".txt", TXT, b""),
    ],
    ids=["pdf", "pptx", "txt", "txt-charset", "txt-empty"],
)
def test_genuine_files_pass(name, ext, mime, content):
    assert content_mismatch(name, ext, mime, content) is None


@pytest.mark.parametrize(
    "name, ext, mime, content, reason",
    [
        ("notes.pdf", ".pdf", PDF, EXE_BYTES, "does not contain PDF data"),
        ("notes.pdf", ".pdf", PDF, b"just some text", "does not contain PDF data"),
        ("slides.pptx", ".pptx", PPTX, _plain_zip(), "does not contain PowerPoint (.pptx) data"),
        ("slides.pptx", ".pptx", PPTX, b"PK\x03\x04 truncated", "does not contain PowerPoint (.pptx) data"),
        ("old.pptx", ".pptx", PPTX, PPT_BYTES, "does not contain PowerPoint (.pptx) data"),
        ("notes.txt", ".txt", TXT, EXE_BYTES, "does not contain text data"),
        ("notes.txt", ".txt", TXT, PDF_BYTES, "does not contain text data"),
        ("notes.txt", ".txt", TXT, b"abc\x00def", "does not contain text data"),
        # Declared type disagrees with the extension, whatever the bytes are.
        ("notes.pdf", ".pdf", TXT, PDF_BYTES, "was sent as text/plain"),
        ("notes.txt", ".txt", "application/x-msdownload", b"hello", "was sent as application/x-msdownload"),
        ("notes.pdf", ".pdf", None, PDF_BYTES, "was sent as an unknown type"),
    ],
    ids=[
        "exe-as-pdf", "text-as-pdf", "zip-as-pptx", "truncated-pptx", "legacy-ppt-renamed-pptx",
        "exe-as-txt", "pdf-as-txt", "nul-in-txt", "pdf-declared-text", "txt-declared-exe",
        "pdf-no-type",
    ],
)
def test_mismatches_are_explained(name, ext, mime, content, reason):
    message = content_mismatch(name, ext, mime, content)
    assert message is not None
    assert reason in message
    assert name in message


# ── Through the upload endpoint ───────────────────────────────────────────────

TEST_DB_URL = "sqlite+aiosqlite:///:memory:"
_engine = create_async_engine(TEST_DB_URL, connect_args={"check_same_thread": False})
_Session = async_sessionmaker(_engine, expire_on_commit=False)


async def _override_db():
    async with _Session() as session:
        try:
            yield session
            await session.commit()
        except Exception:
            await session.rollback()
            raise


@pytest.fixture(autouse=True)
async def setup_db():
    async with _engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    async with _Session() as db:
        db.add(User(email="teacher@demo.com", hashed_password=hash_password("password123"), role="teacher"))
        await db.commit()
    yield
    async with _engine.begin() as conn:
        await conn.run_sync(Base.metadata.drop_all)


@pytest.fixture
async def client():
    app.dependency_overrides[get_db_session] = _override_db
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as ac:
        login = await ac.post("/auth/login", json={"email": "teacher@demo.com", "password": "password123"})
        ac.headers["Authorization"] = f"Bearer {login.json()['access_token']}"
        yield ac
    app.dependency_overrides.clear()


async def _course_count() -> int:
    async with _Session() as db:
        return (await db.execute(select(func.count()).select_from(Course))).scalar_one()


async def test_genuine_pdf_and_pptx_upload_together(client):
    resp = await client.post(
        "/courses/upload",
        data={"name": "Sorting"},
        files=[
            ("files", ("notes.pdf", PDF_BYTES, PDF)),
            ("files", ("slides.pptx", _pptx_bytes(), PPTX)),
        ],
    )
    assert resp.status_code == 200, resp.text


async def test_renamed_executable_is_rejected_before_anything_is_created(client):
    resp = await client.post(
        "/courses/upload",
        data={"name": "Sorting"},
        files={"files": ("notes.pdf", EXE_BYTES, PDF)},
    )
    assert resp.status_code == 400
    assert resp.json()["detail"].startswith("'notes.pdf' does not contain PDF data")
    assert await _course_count() == 0


async def test_declared_type_that_disagrees_with_the_extension_is_rejected(client):
    resp = await client.post(
        "/courses/upload",
        data={"name": "Sorting"},
        files={"files": ("notes.txt", b"plain notes", "application/x-msdownload")},
    )
    assert resp.status_code == 400
    assert "does not match its .txt extension" in resp.json()["detail"]
    assert await _course_count() == 0


@pytest.mark.parametrize(
    "mime", [PPT, PPTX], ids=["declared-ppt", "declared-pptx"]
)
async def test_legacy_ppt_is_no_longer_accepted(client, mime):
    # python-pptx cannot read .ppt, so it used to pass upload and fail in ingestion.
    resp = await client.post(
        "/courses/upload",
        data={"name": "Old deck"},
        files={"files": ("old.ppt", PPT_BYTES, mime)},
    )
    assert resp.status_code == 400
    assert resp.json()["detail"] == "Unsupported file type: old.ppt"
    assert await _course_count() == 0


async def test_one_bad_file_rejects_the_whole_upload(client):
    resp = await client.post(
        "/courses/upload",
        data={"name": "Sorting"},
        files=[
            ("files", ("notes.pdf", PDF_BYTES, PDF)),
            ("files", ("slides.pptx", _plain_zip(), PPTX)),
        ],
    )
    assert resp.status_code == 400
    assert "'slides.pptx'" in resp.json()["detail"]
    assert await _course_count() == 0

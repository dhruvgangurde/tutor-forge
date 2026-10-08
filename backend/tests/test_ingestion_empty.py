"""
tests/test_ingestion_empty.py
-----------------------------
A course with nothing in it fails instead of finishing "ready".

A blank or scanned PDF parsed without error, produced no text, and the course
ended "ready" with 0 chapters -- it looked usable and was empty. Ingestion now
stops with a stored "No readable text" error when the files yield no text, and
a "No chapters" error when the outline comes back empty; both become a plain
sentence on the failed-course card (courses/failure.py).
"""

import io
import json
from unittest.mock import MagicMock

from pptx import Presentation

from agents.ingestion.nodes import (
    NO_CHAPTERS_ERROR,
    NO_READABLE_TEXT_ERROR,
    build_hierarchy_node,
    parse_content_node,
)
from core.config import settings
from courses.failure import failure_reason
from tests.test_ingestion import _state


def _blank_pdf() -> bytes:
    """A valid one-page PDF with no text on the page, like a scan without OCR."""
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, obj in enumerate(objs, 1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % i + obj + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 4\n0000000000 65535 f \n")
    out.write(b"".join(b"%010d 00000 n \n" % o for o in offsets))
    out.write(b"trailer\n<< /Size 4 /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % xref)
    return out.getvalue()


def _blank_deck() -> bytes:
    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[6])  # blank layout: no text frames
    buf = io.BytesIO()
    deck.save(buf)
    return buf.getvalue()


def _file(name: str, content: bytes, mime: str) -> dict:
    return {"filename": name, "content": content, "mime_type": mime}


def test_blank_pdf_fails_with_no_readable_text():
    out = parse_content_node(_state(files=[_file("scan.pdf", _blank_pdf(), "application/pdf")]))
    assert out["status"] == "failed"
    assert out["error"] == NO_READABLE_TEXT_ERROR
    assert out["error"].startswith("No readable text")


def test_blank_deck_and_whitespace_text_also_fail():
    pptx = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    out = parse_content_node(_state(files=[
        _file("empty.pptx", _blank_deck(), pptx),
        _file("spaces.txt", b"   \n\n\t ", "text/plain"),
    ]))
    assert out["status"] == "failed"
    assert out["error"] == NO_READABLE_TEXT_ERROR


def test_a_file_with_text_still_parses():
    out = parse_content_node(_state(files=[_file("notes.txt", b"Merge sort splits arrays.", "text/plain")]))
    assert out.get("status") != "failed"
    assert out["parsed_content"][0]["text"] == "Merge sort splits arrays."


def _parsed() -> list[dict]:
    return [{"source_file": "a.txt", "page_or_slide": None, "text": "some course content"}]


def test_zero_chapters_fails_on_the_default_path():
    gp = MagicMock()
    gp.generate.return_value = json.dumps({"chapters": []})
    out = build_hierarchy_node(_state(parsed_content=_parsed()), gp)
    assert out["status"] == "failed"
    assert out["error"] == NO_CHAPTERS_ERROR


def test_zero_chapters_fails_on_the_ollama_path(monkeypatch):
    monkeypatch.setattr(settings, "llm_provider", "ollama")
    gp = MagicMock()
    gp.generate.return_value = json.dumps({"course_title": "x", "chapters": []})
    out = build_hierarchy_node(_state(parsed_content=_parsed()), gp)
    # The Ollama parser already treats an empty outline as unusable and retries
    # once, so this fails as a parse error before the zero-chapter backstop.
    assert out["status"] == "failed"
    assert out["error"].startswith("Hierarchy JSON parse error")
    assert gp.generate.call_count == 2


def test_the_teacher_sees_a_plain_reason():
    assert failure_reason(NO_READABLE_TEXT_ERROR).startswith("No readable text was found in these files.")
    assert failure_reason(NO_CHAPTERS_ERROR).startswith("No chapters could be found")
